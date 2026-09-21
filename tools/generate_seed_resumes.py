from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "src" / "server" / "Data" / "Seed" / "Resumes"
OUTPUT.mkdir(parents=True, exist_ok=True)


RESUMES = {
    "maya-fernandez-resume.pdf": {
        "name": "Maya Fernandez",
        "title": "Full Stack Engineer",
        "contact": "Chicago, IL | (312) 555-0168 | maya.fernandez@example.com | https://www.linkedin.com/in/maya-fernandez-example",
        "summary": "Full stack engineer with eight years of experience building accessible web products and dependable APIs for growing teams.",
        "skills": "TypeScript, React, Next.js, Node.js, C#, .NET, PostgreSQL, AWS, Docker, Playwright, REST, Git",
        "experience": [
            ("Senior Full Stack Engineer | Northstar Software | 2021 - 2026", [
                "Led development of a React and TypeScript customer portal used by 40,000 monthly users.",
                "Built .NET and PostgreSQL services that reduced median response time by 38 percent.",
                "Added Playwright coverage and GitHub Actions checks, cutting escaped defects by 27 percent.",
            ]),
            ("Software Engineer | Lakeside Labs | 2018 - 2021", [
                "Delivered accessible web workflows with React, Node.js, REST APIs, and AWS.",
                "Partnered with design and support teams to improve task completion for new customers.",
            ]),
        ],
        "education": "Bachelor of Science in Computer Science | University of Illinois Chicago | 2018",
        "certifications": "AWS Certified Developer Associate | 2023",
        "languages": "English, Spanish",
    },
    "ada-whitfield-resume.pdf": {
        "name": "Ada Whitfield",
        "title": "Product Designer",
        "contact": "New York, NY | (212) 555-0184 | ada.whitfield@example.com | https://www.linkedin.com/in/ada-whitfield-example",
        "summary": "Product designer with nine years of experience turning complex workflows into accessible, research-backed products and design systems.",
        "skills": "Product Design, Figma, Design Systems, User Research, Accessibility, Adobe, Agile, Jira, Notion, Testing",
        "experience": [
            ("Senior Product Designer | Harbor Digital | 2020 - 2026", [
                "Owned end-to-end product design for onboarding and account administration workflows.",
                "Created a Figma design system adopted across six product teams and 120 reusable components.",
                "Led user research and accessibility testing that improved task completion by 31 percent.",
            ]),
            ("Product Designer | Civic Studio | 2017 - 2020", [
                "Designed responsive service tools with engineers, researchers, and public-sector partners.",
                "Planned usability studies and translated findings into tested product improvements.",
            ]),
        ],
        "education": "Bachelor of Fine Arts in Communication Design | Pratt Institute | 2017",
        "certifications": "IAAP Certified Professional in Accessibility Core Competencies | 2024",
        "languages": "English, French",
    },
}


ELLIOT = {
    "name": "Elliot Shaw",
    "title": "Backend Engineer",
    "contact": "Denver, CO | (720) 555-0139 | elliot.shaw@example.com | https://www.linkedin.com/in/elliot-shaw-example",
    "summary": "Backend engineer with ten years of experience building secure APIs, event-driven services, and data platforms for high-volume products.",
    "skills": "C#, .NET, PostgreSQL, AWS, Docker, Kubernetes, Kafka, Redis, OpenAPI, REST, Terraform, Datadog",
    "experience": [
        ("Senior Backend Engineer | Summit Systems | 2020 - 2026", [
            "Designed .NET and PostgreSQL services processing more than 12 million requests per day.",
            "Introduced Kafka workflows and Redis caching that reduced peak processing time by 44 percent.",
            "Built AWS, Docker, Kubernetes, Terraform, and Datadog standards for reliable releases.",
        ]),
        ("Software Engineer | Front Range Technology | 2016 - 2020", [
            "Developed C# REST APIs and OpenAPI contracts for billing and account services.",
            "Improved database query performance and added automated integration testing.",
        ]),
    ],
    "education": "Bachelor of Science in Software Engineering | Colorado State University | 2016",
    "certifications": "AWS Certified Solutions Architect Associate | 2022",
    "languages": "English, German",
}


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("ResumeName", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=22, alignment=TA_CENTER, textColor=RGBColor(0, 0, 0), spaceAfter=3),
        "title": ParagraphStyle("ResumeTitle", parent=styles["Normal"], fontName="Helvetica", fontSize=11, leading=14, alignment=TA_CENTER, spaceAfter=3),
        "contact": ParagraphStyle("ResumeContact", parent=styles["Normal"], fontName="Helvetica", fontSize=8.7, leading=11, alignment=TA_CENTER, spaceAfter=10),
        "heading": ParagraphStyle("ResumeHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, alignment=TA_LEFT, textColor=RGBColor(0, 0, 0), spaceBefore=7, spaceAfter=3),
        "body": ParagraphStyle("ResumeBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.4, leading=12, alignment=TA_LEFT, spaceAfter=3),
        "job": ParagraphStyle("ResumeJob", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=9.5, leading=12, alignment=TA_LEFT, spaceBefore=3, spaceAfter=2),
        "bullet": ParagraphStyle("ResumeBullet", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=11.6, leftIndent=12, firstLineIndent=-7, bulletIndent=5, spaceAfter=2),
    }


def create_pdf(path: Path, resume: dict):
    styles = pdf_styles()
    story = [
        Paragraph(resume["name"], styles["name"]),
        Paragraph(resume["title"], styles["title"]),
        Paragraph(resume["contact"], styles["contact"]),
        Paragraph("PROFESSIONAL SUMMARY", styles["heading"]),
        Paragraph(resume["summary"], styles["body"]),
        Paragraph("SKILLS", styles["heading"]),
        Paragraph(resume["skills"], styles["body"]),
        Paragraph("EXPERIENCE", styles["heading"]),
    ]
    for title, bullets in resume["experience"]:
        story.append(Paragraph(title, styles["job"]))
        story.extend(Paragraph(f"• {bullet}", styles["bullet"]) for bullet in bullets)
    story.extend([
        Paragraph("EDUCATION", styles["heading"]),
        Paragraph(resume["education"], styles["body"]),
        Paragraph("CERTIFICATIONS", styles["heading"]),
        Paragraph(resume["certifications"], styles["body"]),
        Paragraph("LANGUAGES", styles["heading"]),
        Paragraph(resume["languages"], styles["body"]),
    ])
    doc = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.7 * inch, leftMargin=0.7 * inch, topMargin=0.55 * inch, bottomMargin=0.55 * inch, title=f"{resume['name']} Resume", author="MyThorneAI ATS Sample Data")
    doc.build(story)


def set_font(run, name="Arial", size=Pt(10), bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = size
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def create_docx(path: Path, resume: dict):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10)

    title_style = doc.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style_element = title_style._element
    title_style_ppr = title_style_element.get_or_add_pPr()
    title_border = title_style_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_style_ppr.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    set_font(title.add_run(resume["name"]), size=Pt(20), bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(resume["title"]), size=Pt(11))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run(resume["contact"]), size=Pt(8.5))

    def heading(value):
        p = doc.add_paragraph(style="Heading 1")
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        set_font(p.add_run(value), size=Pt(10.5), bold=True)

    def body(value, size=9.5, after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.04
        set_font(p.add_run(value), size=Pt(size))
        return p

    heading("PROFESSIONAL SUMMARY")
    body(resume["summary"])
    heading("SKILLS")
    body(resume["skills"])
    heading("EXPERIENCE")
    for job, bullets in resume["experience"]:
        p = body(job, after=1)
        p.paragraph_format.keep_with_next = True
        p.runs[0].bold = True
        for bullet in bullets:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.12)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.02
            set_font(p.add_run(bullet), size=Pt(9.3))
    heading("EDUCATION")
    body(resume["education"])
    heading("CERTIFICATIONS")
    body(resume["certifications"])
    heading("LANGUAGES")
    body(resume["languages"])

    props = doc.core_properties
    props.title = f"{resume['name']} Resume"
    props.author = "MyThorneAI ATS Sample Data"
    props.subject = "Fictional resume parser seed fixture"
    doc.save(path)


for filename, resume in RESUMES.items():
    create_pdf(OUTPUT / filename, resume)
create_docx(OUTPUT / "elliot-shaw-resume.docx", ELLIOT)

print("\n".join(str(path) for path in sorted(OUTPUT.iterdir())))
